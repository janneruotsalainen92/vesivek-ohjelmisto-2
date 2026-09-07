import json

from openpyxl import load_workbook

from vesivek.measure.classify import classify_photo
from vesivek.pipeline import RunRequest, run_measurement
from tests.conftest import write_color_image, write_stick_image


def test_stub_pipeline_writes_png_and_excel(tmp_path):
    photos = tmp_path / "kuvat"
    photos.mkdir()
    write_color_image(photos / "a.jpg", (70, 70, 72))
    write_color_image(photos / "b.jpg", (50, 140, 55))
    write_stick_image(photos / "tikku.jpg", stick_px=180)

    req = RunRequest(
        osoite="Esimerkkitontti",
        kuvat=[photos],
        julkisivu="etela",
        mittatikku_m=1.0,
        kaistan_leveys_m=1.2,
        wfs_mode="stub",
        skip_geocode=True,
        output_dir=tmp_path / "out",
        mittaviivat=True,
        ortho=False,
    )
    result = run_measurement(req)
    assert result.png_path and result.png_path.exists()
    assert result.xlsx_path and result.xlsx_path.exists()
    assert result.geojson_path and result.geojson_path.exists()
    assert result.site.crs == "EPSG:3067"
    assert result.facade.length_m > 1
    length = next(r for r in result.surfaces if r.tyyppi == "julkisivu_pituus")
    assert abs((length.value or 0) - result.facade.length_m) < 0.01
    assert length.unit == "m"
    assert length.ala_m2 is None
    areas = [r for r in result.surfaces if r.kind == "area" and r.value is not None]
    assert areas
    assert result.strip and result.strip.geometry
    assert result.png_path.stat().st_size > 1000
    assert result.xlsx_path.stat().st_size > 1000
    wb = load_workbook(result.xlsx_path)
    assert "WFS-pituus" in wb.sheetnames
    assert "Lineaariset" in wb.sheetnames
    assert "Pinta-alat" in wb.sheetnames


def test_no_width_no_plot_excel_ei_laskettu(tmp_path, rectangle_site):
    gj = tmp_path / "bldg.geojson"
    gj.write_text(
        json.dumps({"type": "FeatureCollection", "features": [rectangle_site.building]}),
        encoding="utf-8",
    )
    photos = tmp_path / "kuvat"
    photos.mkdir()
    write_color_image(photos / "a.jpg", (70, 70, 72))
    req = RunRequest(
        osoite="Testikatu 1",
        kuvat=[photos],
        julkisivu="etela",
        wfs_mode="file",
        geojson=gj,
        output_dir=tmp_path / "out",
        ortho=False,
        mittaviivat=False,
    )
    result = run_measurement(req)
    areas = [r for r in result.surfaces if r.kind == "area"]
    assert areas
    assert all(r.value is None for r in areas)
    wb = load_workbook(result.xlsx_path)
    sheet = wb["Pinta-alat"]
    texts = [str(c.value or "") for row in sheet.iter_rows(min_row=2) for c in row]
    assert any("EI LASKETTU" in t for t in texts)


def test_asphalt_still_classifies(tmp_path):
    path = write_color_image(tmp_path / "asf.jpg", (75, 75, 78), size=(320, 240))
    info = classify_photo(path)
    assert info.fractions["asfaltti"] > 0.5
