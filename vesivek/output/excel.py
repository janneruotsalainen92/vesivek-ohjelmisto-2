from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from vesivek.models import MeasurementResult
from vesivek.valokuva import TYOJARJESTYS, coded_lock_rows, dual_source_verdict, strip_total_lock

HEADER_FILL = PatternFill("solid", fgColor="E85D04")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WARN_FILL = PatternFill("solid", fgColor="FFF3CD")
LOCK_FILL = PatternFill("solid", fgColor="D1FAE5")
UNSURE_FILL = PatternFill("solid", fgColor="FECACA")


def write_excel(result: MeasurementResult, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    _summary(wb.active, result)
    _valokuva_locks(wb.create_sheet("Valokuva-lukot"), result)
    _wfs_length(wb.create_sheet("WFS-pituus"), result)
    _linear(wb.create_sheet("Lineaariset"), result)
    _areas(wb.create_sheet("Pinta-alat"), result)
    _photos(wb.create_sheet("Valokuvat"), result)
    _wfs(wb.create_sheet("WFS"), result)
    if result.qc_ticks:
        _ticks(wb.create_sheet("Mittaviivat"), result)

    wb.save(path)
    return path


def _style_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _autosize(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _summary(ws, result: MeasurementResult) -> None:
    ws.title = "Yhteenveto"
    ws["A1"] = "Vesivek Ohjelma — bot-kokeilu (ei virallinen tuote)"
    ws["A1"].font = Font(bold=True, size=16, color="E85D04")
    ws["A2"] = (
        "Yksi julkisivukaista (seinä → reuna -polygoni). "
        "Excel erottaa: (a) lukittu WFS-pituus, (b) lineaariset m, (c) m² vain kun leveys tunnetaan. "
        "Valokuva-lukot: PRE-LOCK kunnes dual ±10 %; peite → EI VARMENNETTU. Ei salaojaa, ei sadevesiputkia."
    )

    width = result.strip_width_m if result.strip_width_m is not None else "MITTAAMATTA"
    rows = [
        ("Osoite (syöte)", result.site.address_query),
        ("WFS-osuma", result.site.matched_address or "—"),
        ("CRS", result.site.crs),
        ("WFS-lähde", result.site.source_name),
        ("WFS-lukittu", "kyllä" if result.site.verified else "EI — esimerkkigeometria / paikallinen tiedosto"),
        ("Tontti", result.site.plot_source or ("kyllä" if result.site.plot else "ei")),
        ("Orto", result.site.ortho_source.split(" (")[0] if result.site.ortho_source else "ei"),
        ("Kohde-id", result.site.feature_id or "—"),
        ("Julkisivu", result.facade.label_fi),
        ("(a) WFS-pituus m", round(result.facade.length_m, 3)),
        ("Kaistan leveys m", width),
        ("Leveyden lähde", result.strip_width_source),
        ("Kaistan clip", result.strip.clip if result.strip else "—"),
        ("Kaista yhteensä m²", result.strip.area_m2 if result.strip and result.strip.area_m2 is not None else "EI LASKETTU"),
        ("Kaista A m² (tontti)", result.strip.area_m2_plot if result.strip else "—"),
        ("Kaista B m² (puskuri)", result.strip.area_m2_buffer if result.strip else "—"),
        ("Dual", (result.strip.dual_status if result.strip else "—") + " — " + (result.dual_note or "")),
        ("Työvaihe 0–3", result.tyovaihe),
        ("Mittatikku", result.stick.huomio),
        ("Peite / occlusion", result.occlusion_note or ("kyllä" if result.occlusion else "ei")),
        ("Valokuvia", len(result.photos)),
        ("PNG", str(result.png_path) if result.png_path else "—"),
        ("GeoJSON", str(result.geojson_path) if result.geojson_path else "—"),
    ]
    ws.append([])
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
        if k == "WFS-lukittu":
            ws.cell(i, 2).fill = LOCK_FILL if result.site.verified else UNSURE_FILL
        if "EI LASKETTU" in str(v) or v == "MITTAAMATTA":
            ws.cell(i, 2).fill = UNSURE_FILL

    r = 4 + len(rows) + 2
    ws.cell(r, 1, "Huomiot / epävarmuudet").font = Font(bold=True)
    for note in result.warnings + result.site.warnings:
        r += 1
        ws.cell(r, 1, note)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(r, 1).fill = WARN_FILL
        ws.cell(r, 1).alignment = Alignment(wrap_text=True)

    r += 2
    ws.cell(
        r,
        1,
        "Sääntö (Valokuva neliötapa): metrejä ei keksitä. (a) WFS-särmä EPSG:3067 on ainoa lukittu pituus. "
        "(b) Lineaariset m. (c) m² vain kaistapolygonista kun leveys tunnetaan. "
        "Kaksi lähdettä ±10 % = dual_ok; >10 % raportoi A ja B, ei keskiarvoa. "
        "HSV-osuudet = PRE-LOCK, eivät lukittuja metrejä. Peite → EI VARMENNETTU.",
    )
    _autosize(ws, [32, 90, 24, 24])


def _valokuva_locks(ws, result: MeasurementResult) -> None:
    _style_header(ws, ["Tila", "Lukko", "Selite"])
    for status, key, note in coded_lock_rows():
        ws.append([status, key, note])
        last = ws.max_row
        ws.cell(last, 1).fill = LOCK_FILL if status == "KOODATTU" else WARN_FILL
        ws.cell(last, 3).alignment = Alignment(wrap_text=True)
    ws.append([])
    ws.append(["Työvaihe", result.tyovaihe, TYOJARJESTYS[min(result.tyovaihe, 3)][1]])
    ws.append(["Dual", result.strip.dual_status if result.strip else "none", result.dual_note or ""])
    ws.append(
        [
            "CRS",
            result.site.crs,
            "m² vain kun kaista on pinottu WFS-särmään tässä CRS:ssä.",
        ]
    )
    _autosize(ws, [12, 28, 100])


def _wfs_length(ws, result: MeasurementResult) -> None:
    _style_header(ws, ["Kohde", "WFS-pituus_m", "Yksikkö", "Luotettavuus", "Lähde", "Huomio"])
    rec = next((r for r in result.surfaces if r.tyyppi == "julkisivu_pituus"), None)
    ws.append(
        [
            rec.label_fi if rec else result.facade.label_fi,
            round(result.facade.length_m, 3),
            "m",
            "wfs" if result.facade.verified else "esimerkki",
            result.facade.source,
            "Lukittu WFS-geometriaan EPSG:3067. Tätä saraketta ei sekoiteta m²-arvoihin.",
        ]
    )
    ws.cell(2, 4).fill = LOCK_FILL if result.facade.verified else UNSURE_FILL
    _autosize(ws, [40, 16, 10, 16, 36, 80])


def _linear(ws, result: MeasurementResult) -> None:
    _style_header(
        ws,
        ["Kohde", "Tyyppi", "Lineaarinen_m", "Luotettavuus", "Lähde", "Peite", "Huomio"],
    )
    rows = [r for r in result.surfaces if r.kind in {"linear", "wfs"}]
    if not rows:
        ws.append(["(ei lineaarisia rivejä)", "", "", "", "", "", ""])
    for rec in rows:
        val = rec.pituus_m if rec.pituus_m is not None else rec.value
        ws.append(
            [
                rec.label_fi,
                rec.tyyppi,
                val if val is not None else "EI LASKETTU",
                rec.luotettavuus,
                rec.lahde,
                rec.peite or "",
                rec.huomio,
            ]
        )
        last = ws.max_row
        if rec.luotettavuus == "wfs":
            ws.cell(last, 4).fill = LOCK_FILL
        elif rec.luotettavuus in {"epavarma", "esimerkki", "ei_varmennettu"} or val is None:
            ws.cell(last, 3).fill = UNSURE_FILL
            ws.cell(last, 4).fill = UNSURE_FILL
        else:
            ws.cell(last, 4).fill = WARN_FILL
        ws.cell(last, 7).alignment = Alignment(wrap_text=True)
    _autosize(ws, [36, 18, 16, 16, 36, 18, 80])


def _areas(ws, result: MeasurementResult) -> None:
    _style_header(
        ws,
        [
            "Kohde",
            "Tyyppi",
            "ala_m2",
            "ala_m2_B",
            "Luotettavuus",
            "lock_tila",
            "Leveyden lähde",
            "Peite",
            "Osuus (seinästä ulos)",
            "Huomio",
        ],
    )
    rows = [r for r in result.surfaces if r.kind == "area"]
    if result.strip and result.strip.area_m2 is not None:
        luot, lock = strip_total_lock(
            peite=result.occlusion_note or None,
            dual=dual_source_verdict(result.strip.area_m2_plot, result.strip.area_m2_buffer),
            wfs_verified=result.site.verified,
        )
        cell_a = result.strip.area_m2
        if result.strip.dual_status == "disagree":
            cell_a = result.strip.area_m2_plot
        ws.append(
            [
                "Kaista yhteensä (polygoni)",
                "kaista",
                cell_a if cell_a is not None else "EI LASKETTU",
                result.strip.area_m2_buffer if result.strip.dual_status == "disagree" else "",
                luot,
                lock,
                result.strip.width_source,
                result.occlusion_note or "",
                1.0,
                result.strip.huomio,
            ]
        )
    if not rows:
        ws.append(["(ei pinta-aloja)", "", "EI LASKETTU", "epavarma", "", "", "", "Leveys puuttuu tai osuuksia ei ole."])
        ws.cell(ws.max_row, 3).fill = UNSURE_FILL
    for rec in rows:
        val = rec.ala_m2 if rec.ala_m2 is not None else rec.value
        cell_val = val if val is not None else "EI LASKETTU"
        ws.append(
            [
                rec.label_fi,
                rec.tyyppi,
                cell_val,
                rec.ala_m2_b if rec.ala_m2_b is not None else "",
                rec.luotettavuus,
                rec.lock_tila,
                result.strip_width_source,
                rec.peite or "",
                round(rec.share, 4) if rec.share else "",
                rec.huomio,
            ]
        )
        last = ws.max_row
        if rec.luotettavuus in {"epavarma", "esimerkki", "ei_varmennettu", "pre_lock"} or rec.lock_tila in {
            "pre_lock",
            "ei_varmennettu",
            "ei_laskettu",
            "arvio",
        } or val is None:
            ws.cell(last, 3).fill = UNSURE_FILL
            ws.cell(last, 5).fill = UNSURE_FILL
            ws.cell(last, 6).fill = WARN_FILL if rec.lock_tila == "pre_lock" else UNSURE_FILL
        elif rec.luotettavuus == "wfs":
            ws.cell(last, 5).fill = LOCK_FILL
        else:
            ws.cell(last, 5).fill = WARN_FILL
        ws.cell(last, 10).alignment = Alignment(wrap_text=True)
    _autosize(ws, [36, 16, 14, 14, 16, 14, 18, 22, 18, 80])


def _photos(ws, result: MeasurementResult) -> None:
    _style_header(ws, ["Tiedosto", "Px", "Luottamus", "Vallitseva", "Peite", "Huomiot"])
    if not result.photos:
        ws.append(["(ei kuvia)", "", "", "", "", "Pintaosuuksia ei arvioitu."])
    for photo in result.photos:
        top = max(photo.fractions, key=photo.fractions.get)
        ws.append(
            [
                photo.path.name,
                f"{photo.width_px}×{photo.height_px}",
                round(photo.confidence, 3),
                f"{top} {photo.fractions[top]*100:.0f} %",
                photo.occlusion_note if photo.occlusion else "",
                "; ".join(photo.notes),
            ]
        )
    _autosize(ws, [40, 14, 12, 28, 28, 80])


def _wfs(ws, result: MeasurementResult) -> None:
    _style_header(ws, ["Kenttä", "Arvo"])
    site = result.site
    rows = [
        ("Lähde", site.source_name),
        ("URL / kysely", site.source_url or "—"),
        ("CRS", site.crs),
        ("Verified", site.verified),
        ("Feature id", site.feature_id or "—"),
        ("Tontti lähde", site.plot_source or "—"),
        ("Orto lähde", site.ortho_source or "—"),
        ("E", round(site.easting, 3)),
        ("N", round(site.northing, 3)),
        ("Julkisivun särmiä", len(result.facade.edges)),
        ("(a) Julkisivun pituus m", round(result.facade.length_m, 3)),
        ("Kaistan clip", result.strip.clip if result.strip else "—"),
        ("Kaistan leveys m", result.strip_width_m if result.strip_width_m is not None else "MITTAAMATTA"),
    ]
    props = (site.building.get("properties") or {})
    for key, val in list(props.items())[:20]:
        rows.append((f"rakennus.{key}", val))
    if site.plot:
        pprops = site.plot.get("properties") or {}
        for key, val in list(pprops.items())[:12]:
            rows.append((f"tontti.{key}", val))
    for k, v in rows:
        ws.append([k, v])
    _autosize(ws, [28, 100])


def _ticks(ws, result: MeasurementResult) -> None:
    _style_header(ws, ["ID", "Taso", "Tyyppi", "kind", "pituus_m", "Huomio"])
    for t in result.qc_ticks:
        ws.append([t.id, t.tier, t.tyyppi, t.kind, t.length_m, t.huomio])
    ids = [t.id for t in result.qc_ticks]
    ws.append([])
    ws.append(["uniikit ID", "kyllä" if len(ids) == len(set(ids)) else "EI", len(ids), "", "", "Valokuva: MV-* uniikit"])
    _autosize(ws, [16, 8, 16, 14, 14, 70])
